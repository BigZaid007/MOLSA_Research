"""Media-monitoring exports matching the MOLSA رصد الوكالات Word form."""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from datetime import datetime
from typing import Any
from urllib.parse import urlparse

from services.agencies import agency_label, match_agency

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.worksheet import Worksheet

HEADER_GREEN = "92D050"
TITLE_YELLOW = "FFFF00"
THIN_BLACK = Border(
    left=Side(style="thin", color="000000"),
    right=Side(style="thin", color="000000"),
    top=Side(style="thin", color="000000"),
    bottom=Side(style="thin", color="000000"),
)
ARIAL = "Arial"
AR_WEEKDAYS = ["الاثنين", "الثلاثاء", "الاربعاء", "الخميس", "الجمعة", "السبت", "الاحد"]

SOCIAL_SOURCES = {
    "facebook": "فيسبوك",
    "instagram": "رصد ( انستغرام)",
    "x": "منصة X",
    "twitter": "منصة X",
    "tiktok": "تيك توك",
    "youtube": "يوتيوب",
    "telegram": "تليكرام",
    "linkedin": "لينكدإن",
    "reddit": "ريديت",
}

MINISTER_RE = re.compile(
    r"وزير|السيد الوزير|minister of labou?r|minister of social",
    re.I,
)
NEGATIVE_RE = re.compile(
    r"سلب|فساد|أزمة|ازمة|احتجاج|تظاهرة|اعتصام|فضيحة|استياء|انتقاد|scandal|protest|strike|corruption",
    re.I,
)
INTERNATIONAL_RE = re.compile(
    r"دولي|الدولية|الأمم المتحدة|الامم المتحدة|international|united nations|\bilo\b|un\.org",
    re.I,
)

SECTION_SPECS = (
    ("minister", "(اخـــبــار الــســيــد الــوزيــر)", ("ت", "عنوان الخبر", "اسم الموقع")),
    ("agencies", "رصد صدى الاخبار المنشورة في الوكالات والقنوات", ("ت", "عنوان الخبر", "اسم الموقع")),
    ("social", "الاخبار المنشورة في مواقع التواصل", ("ت", "عنوان الخبر", "اسم الموقع")),
    ("subtitles", "السبتايتلات", ("ت", "السبتايتل", "اسم القناة")),
    ("negative", "رصد الاخبار السلبية", ("ت", "عنوان الخبر", "اسم الاعلامي")),
    ("international", "الشأن الدولي", ("ت", "عنوان الخبر", "اسم الموقع")),
)


@dataclass
class ReportSection:
    key: str
    title: str
    headers: tuple[str, str, str]
    rows: list[tuple[str, str, str]] = field(default_factory=list)


@dataclass
class MonitoringReport:
    date_label: str
    sections: list[ReportSection]


def arabic_report_date(when: datetime | None = None) -> str:
    when = when or datetime.now()
    return f"{AR_WEEKDAYS[when.weekday()]}{when.day}/{when.month}/{when.year}"


def report_filename(when: datetime | None = None, ext: str = "docx") -> str:
    when = when or datetime.now()
    return f"رصد_الوكالات_{when.day}-{when.month}-{when.year}.{ext}"


def _blob(item: dict[str, Any]) -> str:
    return " ".join(
        str(item.get(key) or "")
        for key in ("title", "description", "excerpt", "content", "source")
    )


def _title(item: dict[str, Any]) -> str:
    return (item.get("title") or item.get("excerpt") or item.get("description") or "").strip()


def _site_name(item: dict[str, Any]) -> str:
    label = item.get("agency_label") or (item.get("metadata") or {}).get("agency_label")
    if label:
        return str(label)
    matched = agency_label(
        str(item.get("url") or ""),
        str(item.get("source") or ""),
        str(item.get("platform") or ""),
    )
    if matched:
        return matched
    source = str(item.get("source") or "").strip()
    if source and source not in {"agencies", "news", "google", "bing", "rss"}:
        return source
    host = urlparse(str(item.get("url") or "")).netloc.replace("www.", "")
    return host or source


def _is_agency(item: dict[str, Any]) -> bool:
    if (item.get("source") or "").lower() == "agencies":
        return True
    if item.get("agency_label") or (item.get("metadata") or {}).get("agency_id"):
        return True
    return match_agency(str(item.get("url") or ""), str(item.get("source") or ""), str(item.get("platform") or "")) is not None


def _is_social(item: dict[str, Any]) -> bool:
    source = str(item.get("source") or "").lower()
    platform = str(item.get("platform") or "").lower()
    return bool(
        item.get("is_social_mention")
        or source in SOCIAL_SOURCES
        or platform in SOCIAL_SOURCES
    )


def _social_site(item: dict[str, Any]) -> str:
    source = str(item.get("source") or "").lower()
    platform = str(item.get("platform") or "").lower()
    return SOCIAL_SOURCES.get(platform) or SOCIAL_SOURCES.get(source) or _site_name(item)


def _numbered(values: list[tuple[str, str]]) -> list[tuple[str, str, str]]:
    if not values:
        return [("1-", "", "")]
    return [(f"{idx}-", title, site) for idx, (title, site) in enumerate(values, start=1)]


def classify_results(results: list[dict[str, Any]]) -> dict[str, list[tuple[str, str]]]:
    buckets = {key: [] for key, *_ in SECTION_SPECS}
    seen: set[str] = set()

    for item in results:
        title = _title(item)
        if not title:
            continue
        key = (title, item.get("url") or "")
        if key in seen:
            continue
        seen.add(key)
        text = _blob(item)

        if _is_social(item):
            buckets["social"].append((title, _social_site(item)))
        elif _is_agency(item):
            buckets["agencies"].append((title, _site_name(item)))
        elif MINISTER_RE.search(text):
            buckets["minister"].append((title, _site_name(item)))
        elif NEGATIVE_RE.search(text):
            buckets["negative"].append((title, item.get("author") or _site_name(item)))
        elif INTERNATIONAL_RE.search(text):
            buckets["international"].append((title, _site_name(item)))
        else:
            buckets["minister"].append((title, _site_name(item)))

    return buckets


def build_report(payload: dict[str, Any], when: datetime | None = None) -> MonitoringReport:
    buckets = classify_results(list(payload.get("results") or []))
    sections = [
        ReportSection(key, title, headers, _numbered(buckets[key]))
        for key, title, headers in SECTION_SPECS
    ]
    return MonitoringReport(date_label=arabic_report_date(when), sections=sections)


def export_rows(results: list[dict]) -> list[dict]:
    """Flat rows used by tests/callers; same three columns as the Word form."""
    report = build_report({"results": results})
    rows: list[dict] = []
    for section in report.sections:
        for number, title, site in section.rows:
            rows.append(
                {
                    "القسم": section.title,
                    section.headers[0]: number,
                    section.headers[1]: title,
                    section.headers[2]: site,
                }
            )
    return rows


def _style_header_row(sheet: Worksheet, row: int, columns: int) -> None:
    fill = PatternFill("solid", fgColor=HEADER_GREEN)
    font = Font(name=ARIAL, bold=True, size=16, color="000000")
    align = Alignment(horizontal="center", vertical="center", wrap_text=True, readingOrder=2)
    for col in range(1, columns + 1):
        cell = sheet.cell(row, col)
        cell.fill = fill
        cell.font = font
        cell.alignment = align
        cell.border = THIN_BLACK


def _style_data_row(sheet: Worksheet, row: int, columns: int) -> None:
    font = Font(name=ARIAL, size=12)
    align = Alignment(horizontal="center", vertical="center", wrap_text=True, readingOrder=2)
    for col in range(1, columns + 1):
        cell = sheet.cell(row, col)
        cell.font = font
        cell.alignment = align
        cell.border = THIN_BLACK


def build_excel_workbook(payload: dict[str, Any], when: datetime | None = None) -> bytes:
    report = build_report(payload, when)
    wb = Workbook()
    sheet = wb.active
    sheet.title = "رصد الوكالات"
    sheet.sheet_view.rightToLeft = True
    sheet.column_dimensions["A"].width = 8
    sheet.column_dimensions["B"].width = 62
    sheet.column_dimensions["C"].width = 28

    sheet.merge_cells("A1:C1")
    date_cell = sheet["A1"]
    date_cell.value = report.date_label
    date_cell.font = Font(name=ARIAL, bold=True, size=20, underline="single")
    date_cell.fill = PatternFill("solid", fgColor=TITLE_YELLOW)
    date_cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 28

    cursor = 3
    for section in report.sections:
        sheet.merge_cells(start_row=cursor, start_column=1, end_row=cursor, end_column=3)
        title_cell = sheet.cell(cursor, 1, section.title)
        title_cell.font = Font(name=ARIAL, bold=True, size=18)
        title_cell.fill = PatternFill("solid", fgColor=TITLE_YELLOW)
        title_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        sheet.row_dimensions[cursor].height = 24
        cursor += 1

        for col, header in enumerate(section.headers, start=1):
            sheet.cell(cursor, col, header)
        _style_header_row(sheet, cursor, 3)
        sheet.row_dimensions[cursor].height = 28
        cursor += 1

        for number, title, site in section.rows:
            sheet.cell(cursor, 1, number)
            sheet.cell(cursor, 2, title)
            sheet.cell(cursor, 3, site)
            _style_data_row(sheet, cursor, 3)
            sheet.row_dimensions[cursor].height = 36
            cursor += 1
        cursor += 1

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def build_csv_document(payload: dict[str, Any], when: datetime | None = None) -> bytes:
    report = build_report(payload, when)
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow([report.date_label])
    writer.writerow([])
    for section in report.sections:
        writer.writerow([section.title])
        writer.writerow(list(section.headers))
        for row in section.rows:
            writer.writerow(list(row))
        writer.writerow([])
    return buffer.getvalue().encode("utf-8-sig")


def build_word_document(payload: dict[str, Any], when: datetime | None = None) -> bytes:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    report = build_report(payload, when)
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.59)
    section.page_height = Cm(27.94)
    section.left_margin = Cm(1.5)
    section.right_margin = Cm(1.5)
    section.top_margin = Cm(1.4)
    section.bottom_margin = Cm(1.4)
    sect_pr = section._sectPr
    bidi = OxmlElement("w:bidi")
    bidi.set(qn("w:val"), "1")
    sect_pr.append(bidi)

    def _rtl_run(run, *, size: int, bold: bool = False, highlight: str | None = None, underline: bool = False):
        run.font.name = ARIAL
        run.font.size = Pt(size)
        run.bold = bold
        run.font.color.rgb = RGBColor(0, 0, 0)
        run.underline = underline
        rpr = run._element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.append(rfonts)
        rfonts.set(qn("w:ascii"), ARIAL)
        rfonts.set(qn("w:hAnsi"), ARIAL)
        rfonts.set(qn("w:cs"), ARIAL)
        rtl = OxmlElement("w:rtl")
        rpr.append(rtl)
        lang = OxmlElement("w:lang")
        lang.set(qn("w:bidi"), "ar-IQ")
        rpr.append(lang)
        if highlight:
            hl = OxmlElement("w:highlight")
            hl.set(qn("w:val"), highlight)
            rpr.append(hl)

    def _add_heading(text: str, size: int = 18, underline: bool = False):
        para = doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = para.add_run(text)
        _rtl_run(run, size=size, bold=True, highlight="yellow", underline=underline)

    def _shade(cell, fill: str):
        tc_pr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), fill)
        tc_pr.append(shd)

    def _borders(cell):
        tc_pr = cell._tc.get_or_add_tcPr()
        borders = OxmlElement("w:tcBorders")
        for edge in ("top", "left", "bottom", "right"):
            el = OxmlElement(f"w:{edge}")
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), "4")
            el.set(qn("w:space"), "0")
            el.set(qn("w:color"), "000000")
            borders.append(el)
        tc_pr.append(borders)

    def _set_cell(cell, text: str, *, header: bool = False):
        cell.text = ""
        para = cell.paragraphs[0]
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = para.add_run(text)
        _rtl_run(run, size=16 if header else 12, bold=header)
        _borders(cell)
        if header:
            _shade(cell, HEADER_GREEN)

    _add_heading(report.date_label, size=20, underline=True)

    for block in report.sections:
        _add_heading(block.title, size=18)
        table = doc.add_table(rows=1 + len(block.rows), cols=3)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        tbl_pr = table._tbl.tblPr
        if tbl_pr is None:
            tbl_pr = OxmlElement("w:tblPr")
            table._tbl.insert(0, tbl_pr)
        if tbl_pr.find(qn("w:bidiVisual")) is None:
            tbl_pr.append(OxmlElement("w:bidiVisual"))
        widths = (Cm(1.2), Cm(11.0), Cm(6.3))
        for row in table.rows:
            for idx, cell in enumerate(row.cells):
                cell.width = widths[idx]
        for idx, header in enumerate(block.headers):
            _set_cell(table.rows[0].cells[idx], header, header=True)
        for r_idx, values in enumerate(block.rows, start=1):
            for c_idx, value in enumerate(values):
                _set_cell(table.rows[r_idx].cells[c_idx], value)
        doc.add_paragraph()

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def excel_filename(payload: dict[str, Any], stamp: str) -> str:
    return report_filename(ext="xlsx")
